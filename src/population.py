"""District population analysis and forecasting (Mini-Project 1).

This module holds all reusable logic for Mini-Project 1. The notebook
``notebooks/project1_population.ipynb`` imports from here.

Contents
--------
DistrictPopulation
    One district's annual population series, with validation, descriptive
    statistics (``statistics`` and NumPy), growth rates and CAGR.
load_ubos_district
    Reads one district's totals from the UBOS 2022 revised projections workbook.
Forecaster (abstract) and three subclasses
    LinearTrendForecaster, CAGRForecaster, FibonacciRatioForecaster.
mae, rmse, mape, evaluate, select_best
    Hold-out evaluation and model selection.
additional_classrooms
    Planning output for Task 8.
bootstrap_intervals
    Residual-bootstrap prediction intervals (extension).

Units: the supplied series are in thousands of people. UBOS figures are
converted from persons to thousands by ``load_ubos_district``.
"""

from __future__ import annotations

import math
import statistics
from abc import ABC, abstractmethod
from pathlib import Path

import numpy as np
from numpy.typing import ArrayLike, NDArray

__all__ = [
    "DistrictPopulation",
    "load_ubos_district",
    "Forecaster",
    "LinearTrendForecaster",
    "CAGRForecaster",
    "FibonacciRatioForecaster",
    "mae",
    "rmse",
    "mape",
    "evaluate",
    "select_best",
    "additional_classrooms",
    "bootstrap_intervals",
]


# ---------------------------------------------------------------------------
# Shared validation helpers
# ---------------------------------------------------------------------------

def _as_1d_float(values: ArrayLike, label: str) -> NDArray[np.float64]:
    """Convert to a non-empty, finite, 1-D float array or raise ValueError."""
    try:
        arr = np.asarray(values, dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric.") from exc
    if arr.ndim != 1:
        raise ValueError(f"{label} must be one-dimensional, got shape {arr.shape}.")
    if arr.size == 0:
        raise ValueError(f"{label} must not be empty.")
    if not np.all(np.isfinite(arr)):
        raise ValueError(f"{label} must not contain NaN or infinite values.")
    return arr


def _as_years(years: ArrayLike) -> NDArray[np.int64]:
    """Convert to a strictly increasing 1-D integer array of years."""
    arr = _as_1d_float(years, "years")
    if not np.all(arr == np.round(arr)):
        raise ValueError("years must be whole numbers.")
    arr_int = arr.astype(np.int64)
    if arr_int.size > 1 and not np.all(np.diff(arr_int) > 0):
        raise ValueError("years must be strictly increasing.")
    return arr_int


# ---------------------------------------------------------------------------
# Data class
# ---------------------------------------------------------------------------

class DistrictPopulation:
    """Annual population series for one district.

    Parameters
    ----------
    name:
        District name, e.g. ``"Kasese"``.
    years:
        Strictly increasing whole-number years.
    populations:
        Population for each year (non-negative), in ``units``.
    source:
        Where the data came from, e.g. ``"Illustrative (assignment PDF)"``.
    units:
        Unit label for ``populations``. Defaults to ``"thousands"``.

    Raises
    ------
    ValueError
        If the inputs are empty, non-numeric, of unequal length, contain
        NaN/inf or negative populations, or years are not strictly increasing.
    """

    def __init__(
        self,
        name: str,
        years: ArrayLike,
        populations: ArrayLike,
        source: str,
        units: str = "thousands",
    ) -> None:
        if not isinstance(name, str) or not name.strip():
            raise ValueError("name must be a non-empty string.")
        self.name: str = name.strip()
        self.years: NDArray[np.int64] = _as_years(years)
        self.populations: NDArray[np.float64] = _as_1d_float(populations, "populations")
        if self.years.size != self.populations.size:
            raise ValueError(
                f"years ({self.years.size}) and populations "
                f"({self.populations.size}) must have equal length."
            )
        if np.any(self.populations < 0):
            raise ValueError("populations must not be negative.")
        self.source: str = source
        self.units: str = units

    def __repr__(self) -> str:
        return (
            f"DistrictPopulation(name={self.name!r}, "
            f"years={self.years[0]}–{self.years[-1]}, n={len(self)}, "
            f"source={self.source!r})"
        )

    def __len__(self) -> int:
        return int(self.years.size)

    # -- descriptive statistics ------------------------------------------

    def summary(self) -> dict[str, dict[str, float]]:
        """Mean, median, variance and SD computed two ways.

        Returns
        -------
        dict
            ``{"statistics": {...}, "numpy": {...}}``. Each inner dict has
            ``mean``, ``median``, ``sample_variance``, ``population_variance``,
            ``sample_std`` and ``population_std``.

            * ``statistics.variance`` / ``stdev`` use n - 1 (sample).
            * ``statistics.pvariance`` / ``pstdev`` use n (population).
            * NumPy's ``ddof`` sets the divisor to ``n - ddof``: ``ddof=1`` is
              sample, ``ddof=0`` (NumPy's default) is population.

        Raises
        ------
        ValueError
            If fewer than 2 observations (sample variance is undefined).
        """
        if len(self) < 2:
            raise ValueError("At least 2 observations are needed for variance.")
        data = [float(v) for v in self.populations]
        arr = self.populations
        return {
            "statistics": {
                "mean": statistics.mean(data),
                "median": statistics.median(data),
                "sample_variance": statistics.variance(data),
                "population_variance": statistics.pvariance(data),
                "sample_std": statistics.stdev(data),
                "population_std": statistics.pstdev(data),
            },
            "numpy": {
                "mean": float(np.mean(arr)),
                "median": float(np.median(arr)),
                "sample_variance": float(np.var(arr, ddof=1)),
                "population_variance": float(np.var(arr, ddof=0)),
                "sample_std": float(np.std(arr, ddof=1)),
                "population_std": float(np.std(arr, ddof=0)),
            },
        }

    # -- growth --------------------------------------------------------------

    def growth_rates(self) -> NDArray[np.float64]:
        """Year-on-year growth rates ``p[t] / p[t-1] - 1`` (n - 1 values).

        Raises
        ------
        ValueError
            If fewer than 2 observations, or a zero population would cause
            division by zero.
        """
        if len(self) < 2:
            raise ValueError("At least 2 observations are needed for growth rates.")
        if np.any(self.populations[:-1] == 0):
            raise ValueError("Growth rate is undefined when a population is zero.")
        return self.populations[1:] / self.populations[:-1] - 1.0

    def cagr(self) -> float:
        """Compound annual growth rate over the whole series.

        ``(last / first) ** (1 / n_years_elapsed) - 1``, where the number of
        years elapsed is ``years[-1] - years[0]`` (9 for 2015–2024, not 10).

        Raises
        ------
        ValueError
            If fewer than 2 observations or the first population is zero.
        """
        if len(self) < 2:
            raise ValueError("At least 2 observations are needed for CAGR.")
        first, last = self.populations[0], self.populations[-1]
        if first == 0:
            raise ValueError("CAGR is undefined when the first population is zero.")
        elapsed = int(self.years[-1] - self.years[0])
        return float((last / first) ** (1.0 / elapsed) - 1.0)

    # -- train/test split ------------------------------------------------

    def split(self, last_train_year: int) -> tuple[DistrictPopulation, DistrictPopulation]:
        """Split into training (``<= last_train_year``) and test (later) parts.

        Raises
        ------
        ValueError
            If either part would be empty.
        """
        mask = self.years <= last_train_year
        if mask.all() or not mask.any():
            raise ValueError(
                f"last_train_year={last_train_year} leaves an empty train or test set."
            )
        train = DistrictPopulation(
            self.name, self.years[mask], self.populations[mask], self.source, self.units
        )
        test = DistrictPopulation(
            self.name, self.years[~mask], self.populations[~mask], self.source, self.units
        )
        return train, test


# ---------------------------------------------------------------------------
# UBOS loader
# ---------------------------------------------------------------------------

UBOS_SOURCE = (
    "UBOS (2022) Revised Subcounty Population Projections 2015 to 2030 "
    "for 146 Districts (projections, not census counts)"
)


def load_ubos_district(
    path: str | Path,
    name: str,
    start: int = 2015,
    end: int = 2024,
) -> DistrictPopulation:
    """Load one district's total population from the UBOS 2022 workbook.

    The sheet has a two-row header: row 1 holds each year above a
    Male/Female/Total triple, row 2 holds those labels. Each district or city
    has one *total* row (name in column A, column B empty) followed by its
    sub-county rows. Values are converted from persons to thousands.

    Parameters
    ----------
    path:
        Path to ``Revised_Subcounty_population_2015_2030_146_Districts.xlsx``.
    name:
        District name as written in the file, matched case-insensitively
        after stripping spaces (e.g. ``"Kasese"``, ``"Arua City"``).
    start, end:
        Inclusive year range to return.

    Raises
    ------
    FileNotFoundError
        If the workbook does not exist.
    ValueError
        If the district, a year, or a value is missing.
    """
    from openpyxl import load_workbook  # imported here: only this function needs it

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"UBOS workbook not found: {path}")
    if end < start:
        raise ValueError("end must not be before start.")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()

    year_row, label_row = rows[0], rows[1]
    total_col: dict[int, int] = {}
    for col, year in enumerate(year_row):
        if isinstance(year, (int, float)) and col + 2 < len(label_row):
            if str(label_row[col + 2]).strip().lower() == "total":
                total_col[int(year)] = col + 2

    target = name.strip().lower()
    match = None
    for row in rows[2:]:
        if row[0] is not None and row[1] is None and str(row[0]).strip().lower() == target:
            match = row
            break
    if match is None:
        raise ValueError(f"District {name!r} not found in {path.name}.")

    years = list(range(start, end + 1))
    values = []
    for year in years:
        if year not in total_col:
            raise ValueError(f"Year {year} not present in {path.name}.")
        value = match[total_col[year]]
        if value is None:
            raise ValueError(f"Missing {year} value for {name!r}.")
        values.append(float(value) / 1000.0)

    return DistrictPopulation(str(match[0]).strip(), years, values, UBOS_SOURCE)


# ---------------------------------------------------------------------------
# Forecasters
# ---------------------------------------------------------------------------

class Forecaster(ABC):
    """Abstract base class for annual forecasting models.

    Subclasses implement ``_fit`` and ``_predict``. The base class handles
    validation so it is not repeated in each model. Years must be
    consecutive (annual steps) because every model forecasts one year at a time.
    """

    def __init__(self) -> None:
        self._fitted: bool = False
        self.years_: NDArray[np.int64] | None = None
        self.values_: NDArray[np.float64] | None = None

    @property
    def name(self) -> str:
        """Readable model name for tables and legends."""
        return type(self).__name__.replace("Forecaster", "")

    def fit(self, years: ArrayLike, values: ArrayLike) -> Forecaster:
        """Fit the model and return ``self`` (so calls can be chained).

        Raises
        ------
        ValueError
            If inputs are invalid, fewer than 2 points, or years are not consecutive.
        """
        yrs = _as_years(years)
        vals = _as_1d_float(values, "values")
        if yrs.size != vals.size:
            raise ValueError("years and values must have equal length.")
        if yrs.size < 2:
            raise ValueError("At least 2 observations are needed to fit.")
        if not np.all(np.diff(yrs) == 1):
            raise ValueError("years must be consecutive (annual steps).")
        self.years_, self.values_ = yrs, vals
        self._fit(yrs, vals)
        self._fitted = True
        return self

    def predict(self, horizon: int) -> NDArray[np.float64]:
        """Forecast the next ``horizon`` years after the last fitted year.

        Raises
        ------
        RuntimeError
            If called before ``fit``.
        ValueError
            If ``horizon`` is not a positive integer.
        """
        if not self._fitted:
            raise RuntimeError(f"{self.name}: call fit() before predict().")
        if isinstance(horizon, bool) or not isinstance(horizon, (int, np.integer)) or horizon < 1:
            raise ValueError("horizon must be a positive integer.")
        return self._predict(int(horizon))

    def fitted_values(self) -> NDArray[np.float64] | None:
        """In-sample fitted values, or ``None`` if the model has none."""
        if not self._fitted:
            raise RuntimeError(f"{self.name}: call fit() first.")
        return self._fitted_values()

    def forecast_years(self, horizon: int) -> NDArray[np.int64]:
        """Calendar years matching ``predict(horizon)``."""
        if not self._fitted or self.years_ is None:
            raise RuntimeError(f"{self.name}: call fit() first.")
        return self.years_[-1] + np.arange(1, horizon + 1)

    @abstractmethod
    def _fit(self, years: NDArray[np.int64], values: NDArray[np.float64]) -> None:
        """Learn model parameters from validated data."""

    @abstractmethod
    def _predict(self, horizon: int) -> NDArray[np.float64]:
        """Return ``horizon`` forecasts."""

    @abstractmethod
    def _fitted_values(self) -> NDArray[np.float64] | None:
        """Return in-sample fitted values or ``None``."""


class LinearTrendForecaster(Forecaster):
    """Straight-line trend ``y = intercept + slope * (year - base_year)``.

    Fitted with ``np.polyfit`` (degree 1, least squares). Years are centred
    on the first year to keep the fit numerically well conditioned; the
    predictions are identical to fitting on raw years.
    """

    def _fit(self, years: NDArray[np.int64], values: NDArray[np.float64]) -> None:
        self.base_year_: int = int(years[0])
        slope, intercept = np.polyfit(years - self.base_year_, values, deg=1)
        self.slope_: float = float(slope)
        self.intercept_: float = float(intercept)

    def _line(self, years: NDArray[np.int64]) -> NDArray[np.float64]:
        return self.intercept_ + self.slope_ * (years - self.base_year_)

    def _predict(self, horizon: int) -> NDArray[np.float64]:
        return self._line(self.forecast_years(horizon))

    def _fitted_values(self) -> NDArray[np.float64]:
        assert self.years_ is not None
        return self._line(self.years_)


class CAGRForecaster(Forecaster):
    """Constant-percentage growth at the training-period CAGR.

    ``g = (last / first) ** (1 / (n - 1)) - 1`` and
    ``forecast[h] = last * (1 + g) ** h``. Because ``g`` uses only the two
    endpoints, the fitted curve passes exactly through the first and last points.
    """

    def _fit(self, years: NDArray[np.int64], values: NDArray[np.float64]) -> None:
        if values[0] <= 0:
            raise ValueError("CAGR model needs a positive first value.")
        self.growth_rate_: float = float(
            (values[-1] / values[0]) ** (1.0 / (len(values) - 1)) - 1.0
        )

    def _predict(self, horizon: int) -> NDArray[np.float64]:
        assert self.values_ is not None
        steps = np.arange(1, horizon + 1)
        return self.values_[-1] * (1.0 + self.growth_rate_) ** steps

    def _fitted_values(self) -> NDArray[np.float64]:
        assert self.values_ is not None
        steps = np.arange(len(self.values_))
        return self.values_[0] * (1.0 + self.growth_rate_) ** steps


def fibonacci(n: int) -> NDArray[np.int64]:
    """First ``n`` Fibonacci numbers starting ``1, 1, 2, 3, ...``."""
    if n < 1:
        raise ValueError("n must be at least 1.")
    seq = [1, 1]
    while len(seq) < n:
        seq.append(seq[-1] + seq[-2])
    return np.array(seq[:n], dtype=np.int64)


class FibonacciRatioForecaster(Forecaster):
    """Previous cohort's model: scale the last value by successive Fibonacci ratios.

    With ``F = 1, 1, 2, 3, 5, ...`` the ratios are ``r_k = F[k+1] / F[k]``
    (1, 2, 1.5, 1.667, 1.6, ... -> 1.618). Stated assumption: ratios
    *compound*, starting from ``r_start``:
    ``forecast[h] = last * r_start * r_(start+1) * ... * r_(start+h-1)``.

    The model learns nothing from the data except the last value, so it has
    no in-sample fit (``fitted_values()`` returns ``None``).

    Parameters
    ----------
    start_index:
        Which ratio to begin with (1 means ``F2/F1 = 1``).
    """

    def __init__(self, start_index: int = 1) -> None:
        super().__init__()
        if isinstance(start_index, bool) or not isinstance(start_index, int) or start_index < 1:
            raise ValueError("start_index must be a positive integer.")
        self.start_index = start_index

    def ratios(self, horizon: int) -> NDArray[np.float64]:
        """The ``horizon`` Fibonacci ratios used for forecasting."""
        fib = fibonacci(self.start_index + horizon + 1).astype(float)
        k = np.arange(self.start_index, self.start_index + horizon)  # 1-based ratio index
        return fib[k] / fib[k - 1]

    def _fit(self, years: NDArray[np.int64], values: NDArray[np.float64]) -> None:
        pass  # nothing to learn: only the last value (stored by the base class) is used

    def _predict(self, horizon: int) -> NDArray[np.float64]:
        assert self.values_ is not None
        return self.values_[-1] * np.cumprod(self.ratios(horizon))

    def _fitted_values(self) -> None:
        return None


# ---------------------------------------------------------------------------
# Evaluation
# ---------------------------------------------------------------------------

def _paired(actual: ArrayLike, predicted: ArrayLike) -> tuple[NDArray[np.float64], NDArray[np.float64]]:
    a = _as_1d_float(actual, "actual")
    p = _as_1d_float(predicted, "predicted")
    if a.shape != p.shape:
        raise ValueError(f"actual {a.shape} and predicted {p.shape} must have the same shape.")
    return a, p


def mae(actual: ArrayLike, predicted: ArrayLike) -> float:
    """Mean absolute error, in the data's units."""
    a, p = _paired(actual, predicted)
    return float(np.mean(np.abs(a - p)))


def rmse(actual: ArrayLike, predicted: ArrayLike) -> float:
    """Root mean squared error, in the data's units (penalises large misses)."""
    a, p = _paired(actual, predicted)
    return float(np.sqrt(np.mean((a - p) ** 2)))


def mape(actual: ArrayLike, predicted: ArrayLike) -> float:
    """Mean absolute percentage error, in percent.

    Raises
    ------
    ValueError
        If any actual value is zero (percentage error undefined).
    """
    a, p = _paired(actual, predicted)
    if np.any(a == 0):
        raise ValueError("MAPE is undefined when an actual value is zero.")
    return float(100.0 * np.mean(np.abs((a - p) / a)))


def evaluate(
    district: DistrictPopulation,
    models: list[Forecaster],
    last_train_year: int = 2021,
) -> list[dict[str, object]]:
    """Fit each model on the training years and score it on the test years.

    Note: each model in ``models`` is (re)fitted in place.

    Returns
    -------
    list of dict
        One row per model: ``model``, ``MAE``, ``RMSE``, ``MAPE`` and
        ``predictions`` (test-period forecasts).
    """
    if not models:
        raise ValueError("models must not be empty.")
    train, test = district.split(last_train_year)
    rows: list[dict[str, object]] = []
    for model in models:
        model.fit(train.years, train.populations)
        pred = model.predict(len(test))
        rows.append(
            {
                "model": model.name,
                "MAE": mae(test.populations, pred),
                "RMSE": rmse(test.populations, pred),
                "MAPE": mape(test.populations, pred),
                "predictions": pred,
            }
        )
    return rows


def select_best(rows: list[dict[str, object]]) -> str:
    """Name of the model with the lowest RMSE (MAPE breaks ties)."""
    if not rows:
        raise ValueError("rows must not be empty.")
    best = min(rows, key=lambda r: (float(r["RMSE"]), float(r["MAPE"])))  # type: ignore[arg-type]
    return str(best["model"])


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------

def additional_classrooms(
    pop_now: float,
    pop_future: float,
    school_age_share: float = 0.18,
    pupils_per_classroom: int = 53,
    units_multiplier: float = 1000.0,
) -> int:
    """Extra classrooms needed for population growth between two dates.

    ``ceil(share * max(0, pop_future - pop_now) * units_multiplier / pupils)``.
    Assumes current capacity exactly meets current need; a declining
    population needs 0 additional classrooms (never a negative number).

    Raises
    ------
    ValueError
        If populations are negative/non-finite, the share is outside (0, 1],
        or pupils_per_classroom is not a positive integer.
    """
    for label, value in (("pop_now", pop_now), ("pop_future", pop_future)):
        if not math.isfinite(value) or value < 0:
            raise ValueError(f"{label} must be a finite, non-negative number.")
    if not 0 < school_age_share <= 1:
        raise ValueError("school_age_share must be in (0, 1].")
    if isinstance(pupils_per_classroom, bool) or not isinstance(pupils_per_classroom, int) or pupils_per_classroom < 1:
        raise ValueError("pupils_per_classroom must be a positive integer.")
    extra_pupils = school_age_share * max(0.0, pop_future - pop_now) * units_multiplier
    # round first so floating-point noise (e.g. 53.0000000001) does not add a classroom
    return int(math.ceil(round(extra_pupils / pupils_per_classroom, 9)))


# ---------------------------------------------------------------------------
# Extension: residual bootstrap
# ---------------------------------------------------------------------------

def bootstrap_intervals(
    model_cls: type[Forecaster],
    years: ArrayLike,
    values: ArrayLike,
    horizon: int,
    n_boot: int = 1000,
    level: float = 0.95,
    seed: int = 2026,
) -> tuple[NDArray[np.float64], NDArray[np.float64], NDArray[np.float64]]:
    """Residual-bootstrap prediction intervals.

    1. Fit the model; residuals = actual - fitted.
    2. For each of ``n_boot`` resamples: build a pseudo-series
       (fitted + residuals drawn with replacement), refit a *fresh* model,
       forecast, and add one resampled residual per future year for noise.
    3. Take the ``(1 - level)/2`` and ``(1 + level)/2`` percentiles.

    Assumes residuals are independent and identically distributed, which
    trend residuals often are not; with few residuals the intervals are
    likely too narrow.

    Returns
    -------
    (point, lower, upper)
        Point forecast from the original fit and the interval bounds.

    Raises
    ------
    ValueError
        If the model has no in-sample fit (e.g. Fibonacci), or arguments are invalid.
    """
    if n_boot < 1:
        raise ValueError("n_boot must be at least 1.")
    if not 0 < level < 1:
        raise ValueError("level must be between 0 and 1.")
    base = model_cls().fit(years, values)
    fitted = base.fitted_values()
    if fitted is None:
        raise ValueError(f"{base.name} has no in-sample fit, so residuals are undefined.")
    assert base.years_ is not None and base.values_ is not None
    residuals = base.values_ - fitted
    point = base.predict(horizon)

    rng = np.random.default_rng(seed)
    sims = np.empty((n_boot, horizon))
    for b in range(n_boot):
        pseudo = fitted + rng.choice(residuals, size=residuals.size, replace=True)
        model = model_cls().fit(base.years_, pseudo)
        sims[b] = model.predict(horizon) + rng.choice(residuals, size=horizon, replace=True)

    alpha = (1.0 - level) / 2.0
    lower = np.percentile(sims, 100 * alpha, axis=0)
    upper = np.percentile(sims, 100 * (1 - alpha), axis=0)
    return point, lower, upper
